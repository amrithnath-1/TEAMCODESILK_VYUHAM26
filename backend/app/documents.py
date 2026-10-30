"""Document intelligence: validation, text extraction, chunking, indexing and retrieval.

Retrieval choice: SQLite FTS5 with BM25 ranking, scoped per project. No embedding provider is used
(Claude's API does not provide embeddings); lexical retrieval is sufficient for per-project document sets.
All extracted text is untrusted: suspicious passages are flagged and always passed to the model as data.
"""
import csv
import hashlib
import io
import re
import zipfile
from pathlib import Path
from typing import Dict, List, Optional, Tuple

from fastapi import HTTPException

from . import db
from .audit import utcnow

ALLOWED = {".pdf", ".docx", ".pptx", ".txt", ".csv", ".xlsx"}
RESUME_ALLOWED = {".pdf", ".docx", ".txt"}
MAX_BYTES = 20 * 1024 * 1024
MAX_UNZIPPED = 200 * 1024 * 1024
CHUNK_CHARS = 1000
CHUNK_OVERLAP = 150

INJECTION_RE = re.compile(
    r"ignore (all |any )?(the )?(previous|prior|above) (instructions|prompts?)|disregard (the )?(system|previous)|"
    r"you are now|system prompt|reveal (your|the) (instructions|prompt|api key)|act as (an? )?(admin|administrator)|"
    r"assign all tasks to", re.I)
STOPWORDS = set("a an and are as at be but by can do does for from has have how i in is it its me my of on or our should "
                "that the their this to was we what when where which who why will with would you your about please tell".split())


def validate_upload(filename: str, data: bytes, allowed=ALLOWED) -> str:
    name = Path(filename or "").name
    ext = Path(name).suffix.lower()
    if not name or ext not in allowed:
        raise ValueError(f"Unsupported file type '{ext or 'none'}'. Allowed: {', '.join(sorted(allowed))}")
    if not data:
        raise ValueError("The file is empty")
    if len(data) > MAX_BYTES:
        raise ValueError(f"File is {len(data) // (1024 * 1024)} MB; the limit is {MAX_BYTES // (1024 * 1024)} MB")
    if ext == ".pdf" and not data[:1024].lstrip().startswith(b"%PDF"):
        raise ValueError("This file does not look like a real PDF")
    if ext in (".docx", ".pptx", ".xlsx"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                if sum(i.file_size for i in z.infolist()) > MAX_UNZIPPED:
                    raise ValueError("The archive expands to an unsafe size")
        except zipfile.BadZipFile:
            raise ValueError(f"This file is not a valid {ext} document")
    return ext


def _decode(data: bytes) -> str:
    for enc in ("utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(enc)
        except (UnicodeDecodeError, UnicodeError):
            continue
    return data.decode("utf-8", errors="replace")


def extract(ext: str, data: bytes) -> List[Tuple[str, str]]:
    """Returns (location, text) sections; location is a page, slide, sheet or section reference."""
    if ext == ".txt":
        return [("Text", _decode(data))]
    if ext == ".csv":
        rows = list(csv.reader(io.StringIO(_decode(data))))
        return [("CSV rows", "\n".join(" | ".join(r) for r in rows[:5000]))]
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return [(f"Page {i + 1}", p.extract_text() or "") for i, p in enumerate(reader.pages)]
    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        out: List[Tuple[str, str]] = []
        section, buf = "Document", []
        for para in d.paragraphs:
            if para.style is not None and para.style.name.startswith("Heading") and para.text.strip():
                if buf:
                    out.append((section, "\n".join(buf)))
                section, buf = f"Section: {para.text.strip()[:80]}", []
            elif para.text.strip():
                buf.append(para.text.strip())
        if buf:
            out.append((section, "\n".join(buf)))
        for i, table in enumerate(d.tables):
            out.append((f"Table {i + 1}", "\n".join(" | ".join(c.text.strip() for c in r.cells) for r in table.rows)))
        return out
    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        out = []
        for i, slide in enumerate(prs.slides):
            texts = [s.text_frame.text.strip() for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
            for s in slide.shapes:
                if s.has_table:
                    texts.extend(" | ".join(c.text.strip() for c in r.cells) for r in s.table.rows)
            out.append((f"Slide {i + 1}", "\n".join(texts)))
        return out
    if ext == ".xlsx":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        out = []
        for ws in wb.worksheets:
            lines = []
            for row in ws.iter_rows(values_only=True, max_row=5000):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    lines.append(" | ".join(cells))
            out.append((f"Sheet: {ws.title}", "\n".join(lines)))
        return out
    raise ValueError(f"No extractor for {ext}")


def chunk_sections(sections: List[Tuple[str, str]]) -> List[Tuple[str, str]]:
    chunks: List[Tuple[str, str]] = []
    for location, text in sections:
        text = re.sub(r"[ \t]+", " ", text).strip()
        start = 0
        while start < len(text):
            end = min(len(text), start + CHUNK_CHARS)
            if end < len(text):   # prefer a sentence or line boundary
                cut = max(text.rfind("\n", start, end), text.rfind(". ", start, end))
                if cut > start + CHUNK_CHARS // 2:
                    end = cut + 1
            piece = text[start:end].strip()
            if piece:
                chunks.append((location, piece))
            if end >= len(text):
                break
            start = end - CHUNK_OVERLAP
    return chunks


# ------------------------------------------------------------------ lifecycle
def store_upload(project_id: str, filename: str, data: bytes, username: str) -> Dict:
    """Validates and records an upload. Returns the document row; the caller schedules process()."""
    ext = validate_upload(filename, data)
    name = Path(filename).name
    sha = hashlib.sha256(data).hexdigest()
    dup = db.one("SELECT * FROM documents WHERE project_id = ? AND sha256 = ? AND is_current = 1", (project_id, sha))
    if dup:
        return {**dup, "duplicate": True}
    prev = db.one("SELECT MAX(version) AS v FROM documents WHERE project_id = ? AND filename = ?", (project_id, name))
    version = (prev["v"] or 0) + 1
    doc_id = db.new_id("DOC")
    path = db.UPLOAD_DIR / f"{doc_id}{ext}"
    path.write_bytes(data)
    with db.tx() as c:
        c.execute("INSERT INTO documents (id, org_id, project_id, filename, version, sha256, size, ext, status, stored_path, uploaded_by, uploaded_at) "
                  "VALUES (?,?,?,?,?,?,?,?,?,?,?,?)",
                  (doc_id, db.ORG_ID, project_id, name, version, sha, len(data), ext, "uploaded", str(path), username, utcnow()))
    return db.one("SELECT * FROM documents WHERE id = ?", (doc_id,))


def _remove_chunks(c, document_id: str) -> None:
    ids = [r[0] for r in c.execute("SELECT id FROM chunks WHERE document_id = ?", (document_id,)).fetchall()]
    c.executemany("DELETE FROM chunks_fts WHERE chunk_id = ?", [(i,) for i in ids])
    c.execute("DELETE FROM chunks WHERE document_id = ?", (document_id,))


def process(document_id: str) -> None:
    """Extract, chunk and index one document. Never raises: failures are recorded on the document."""
    doc = db.one("SELECT * FROM documents WHERE id = ?", (document_id,))
    if not doc:
        return
    try:
        with db.tx() as c:
            c.execute("UPDATE documents SET status = 'extracting', error = NULL WHERE id = ?", (document_id,))
        sections = extract(doc["ext"], Path(doc["stored_path"]).read_bytes())
        chunks = chunk_sections(sections)
        if not chunks:
            raise ValueError("No readable text found. A scanned PDF needs OCR before it can be indexed.")
        suspicious = 0
        with db.tx() as c:
            _remove_chunks(c, document_id)
            for seq, (location, text) in enumerate(chunks):
                flag = 1 if INJECTION_RE.search(text) else 0
                suspicious += flag
                chunk_id = db.new_id("CHK")
                c.execute("INSERT INTO chunks (id, document_id, project_id, seq, location, text, suspicious) VALUES (?,?,?,?,?,?,?)",
                          (chunk_id, document_id, doc["project_id"], seq, location, text, flag))
                c.execute("INSERT INTO chunks_fts (text, chunk_id, project_id) VALUES (?,?,?)", (text, chunk_id, doc["project_id"]))
            # A new version supersedes older versions of the same file.
            older = [r[0] for r in c.execute("SELECT id FROM documents WHERE project_id = ? AND filename = ? AND version < ? AND is_current = 1",
                                             (doc["project_id"], doc["filename"], doc["version"])).fetchall()]
            for old in older:
                _remove_chunks(c, old)
                c.execute("UPDATE documents SET is_current = 0, status = 'superseded', chunk_count = 0 WHERE id = ?", (old,))
            c.execute("UPDATE documents SET status = 'indexed', chunk_count = ?, suspicious_count = ?, error = NULL WHERE id = ?",
                      (len(chunks), suspicious, document_id))
    except Exception as exc:   # corrupt or unreadable files must surface as a status, not a crash
        with db.tx() as c:
            c.execute("UPDATE documents SET status = 'failed', error = ? WHERE id = ?", (str(exc)[:400], document_id))


def list_documents(project_id: str) -> List[Dict]:
    return db.rows("SELECT id, project_id, filename, version, size, ext, status, error, is_current, chunk_count, suspicious_count, "
                   "uploaded_by, uploaded_at FROM documents WHERE project_id = ? ORDER BY is_current DESC, uploaded_at DESC", (project_id,))


def get_document(project_id: str, document_id: str) -> Dict:
    doc = db.one("SELECT * FROM documents WHERE id = ? AND project_id = ?", (document_id, project_id))
    if not doc:
        raise HTTPException(404, "Document not found in this project")
    return doc


def delete_document(project_id: str, document_id: str) -> None:
    doc = get_document(project_id, document_id)
    with db.tx() as c:
        _remove_chunks(c, document_id)
        c.execute("DELETE FROM documents WHERE id = ?", (document_id,))
    Path(doc["stored_path"]).unlink(missing_ok=True)


def preview(project_id: str, document_id: str, limit: int = 6) -> Dict:
    doc = get_document(project_id, document_id)
    chunks = db.rows("SELECT location, text, suspicious FROM chunks WHERE document_id = ? ORDER BY seq LIMIT ?", (document_id, limit))
    return {"filename": doc["filename"], "version": doc["version"], "status": doc["status"], "chunks": chunks}


def _fts_query(query: str) -> Optional[str]:
    words = [w for w in re.findall(r"[A-Za-z0-9]{2,}", query.lower()) if w not in STOPWORDS]
    return " OR ".join(f'"{w}"' for w in words[:20]) or None


def search(project_id: str, query: str, limit: int = 6) -> List[Dict]:
    """BM25 search restricted to one project's current documents (prevents cross-project leakage)."""
    q = _fts_query(query)
    if not q:
        return []
    return db.rows(
        "SELECT c.id, c.text, c.location, c.suspicious, d.filename, d.version, bm25(chunks_fts) AS rank "
        "FROM chunks_fts JOIN chunks c ON c.id = chunks_fts.chunk_id JOIN documents d ON d.id = c.document_id "
        "WHERE chunks_fts MATCH ? AND chunks_fts.project_id = ? AND c.project_id = ? AND d.is_current = 1 "
        "ORDER BY rank LIMIT ?", (q, project_id, project_id, limit))


def project_text(project_id: str, max_chars: int = 150_000) -> Tuple[List[Dict], bool]:
    """All current chunks for a project, up to the context limit. The flag reports truncation."""
    rows = db.rows("SELECT c.text, c.location, c.suspicious, d.filename FROM chunks c JOIN documents d ON d.id = c.document_id "
                   "WHERE c.project_id = ? AND d.is_current = 1 AND d.status = 'indexed' ORDER BY d.uploaded_at, c.seq", (project_id,))
    out, total = [], 0
    for r in rows:
        total += len(r["text"])
        if total > max_chars:
            return out, True
        out.append(r)
    return out, False


def as_untrusted_block(chunks: List[Dict]) -> str:
    parts = []
    for c in chunks:
        warn = ' flagged="possible-prompt-injection"' if c.get("suspicious") else ""
        parts.append(f'<untrusted_document source="{c["filename"]}" location="{c["location"]}"{warn}>\n{c["text"]}\n</untrusted_document>')
    return "\n".join(parts)
